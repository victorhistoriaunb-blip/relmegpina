"""Testes do Clipping de Novas Proposições + Filtro Inteligente Family Talks."""
import pytest
from docx import Document

import family_talks
from exportador_local import (
    ClippingError,
    _caminho_modelo,
    _data_br,
    _substituir_corpo,
)
from family_talks import filtrar

CAMARA_ITENS = [
    {
        "siglaTipo": "PL",
        "numero": "595",
        "ano": 2026,
        "id": 42,
        "ementa": "Altera a CLT para ampliar a licença maternidade de 120 para 180 dias.",
        "autor": "Fulano de Tal",
        "data": "2026-09-01T00:00:00",
    },
    {
        "siglaTipo": "PL",
        "numero": "777",
        "ano": 2026,
        "id": 43,
        "ementa": "Dispõe sobre alienação parental e guarda compartilhada.",
        "autor": "Ciclana So e Tal",
        "data": "2026-09-02",
    },
]
SENADO_ITENS = [
    {
        "sigla": "PL",
        "numero": "101",
        "ano": 2026,
        "codigo": 7,
        "ementa": "Institui política de atenção integral à primeira infância.",
        "autor": "Senador Beltrano",
        "data": "2026-09-03",
    },
    {
        "sigla": "PL",
        "numero": "202",
        "ano": 2026,
        "codigo": 8,
        "ementa": "Trata da reforma tributária e dos impostos sobre consumo.",
        "autor": "Senadora Beltrana",
        "data": "2026-09-04",
    },
]


def _doc_com_corpo():
    doc = Document()
    doc.add_paragraph("Plano de Ação", style=doc.styles["Title"])
    doc.add_paragraph("Câmara dos Deputados", style=doc.styles["Title"])
    doc.add_paragraph("PL 595/2024")
    doc.add_paragraph("ementa: parametro")
    doc.add_paragraph("autores: parametro")
    doc.add_paragraph("data: parametro")
    doc.add_paragraph("-- marcador --")
    doc.add_paragraph("Senado Federal", style=doc.styles["Title"])
    doc.add_paragraph("Outro texto fixo")
    return doc


def test_caminho_modelo_ausente():
    try:
        model_path = _caminho_modelo()
    except ClippingError as exc:
        assert "MODELO A SER SEGUIDO" in str(exc)
        pytest.skip("Sem modelo na pasta de entregas — erro orientativo validado")
    assert model_path.name == "MODELO A SER SEGUIDO.docx"


def test_data_br():
    assert _data_br("2026-09-07T00:00:00") == "07/09/2026"
    assert _data_br("2026-09-07") == "07/09/2026"
    assert _data_br(None) is None
    assert _data_br("") is None


def test_filtro_family_talks_escopo():
    aprovados = filtrar(CAMARA_ITENS + SENADO_ITENS)
    ementas = [i["ementa"] for i in aprovados]
    assert len(aprovados) == 2
    assert any("licença maternidade" in e for e in ementas)  # lei: vínculo e parentalidade
    assert any("primeira infância" in e for e in ementas)  # Lei: primeira infância
    assert not any("alienação parental" in e for e in ementas)  # exclusão automática
    assert not any("reforma tributária" in e for e in ementas)  # pauta estranha


def test_substituir_corpo_formata_blocos():
    doc = _doc_com_corpo()
    _substituir_corpo(doc, CAMARA_ITENS, SENADO_ITENS)
    textos = [p.text for p in doc.paragraphs]

    assert "Câmara dos Deputados" in textos
    assert "Senado Federal" in textos
    assert any("PL 595/2026" in t for t in textos)
    assert any("PL 101/2026" in t for t in textos)
    assert any("01/09/2026" in t for t in textos)  # data DD/MM/AAAA da Câmara
    assert any("03/09/2026" in t for t in textos)  # data DD/MM/AAAA do Senado
    assert any("Ciclana So e Tal" in t for t in textos)
    assert not any("-- marcador --" in t for t in textos)
    assert not any("parametro" in t for t in textos)  # placeholders substituídos

    alvos = [
        rel.target_ref
        for rel in doc.part.rels.values()
        if getattr(rel, "is_external", False)
    ]
    assert any("fichadetramitacao?idProposicao=42" in a for a in alvos)
    assert any("/materia/7" in a for a in alvos)


def test_matriz_family_talks_temas_e_exclusoes():
    assert family_talks._TEMAS_PRIORITARIOS  # carregada do backend
    assert family_talks._TEMAS_FORA_ESCOPO  # impedimentos ativos
    aprovados = filtrar(CAMARA_ITENS + SENADO_ITENS)
    assert len(family_talks.temas_detectados(aprovados)) >= 2


def test_rota_clipping_rejeita_template_ausente():
    from fastapi.testclient import TestClient

    import main

    with TestClient(main.app) as cliente:
        resposta = cliente.post(
            "/api/exportar/clipping-semanal",
            params={"keywords": "infância, idosos"},
        )
    assert resposta.status_code == 400
    assert "MODELO A SER SEGUIDO" in resposta.json()["detail"]